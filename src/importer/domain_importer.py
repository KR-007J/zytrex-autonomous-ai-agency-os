"""Domain Importer supporting CSV, XLSX, JSON, TXT with validation, deduplication, and verification enqueueing."""

from __future__ import annotations
import io
import csv
import json
import re
import asyncio
from typing import List, Dict, Any, Optional, Set, Union
from urllib.parse import urlparse
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from src.database.db import get_db
from src.database.models import Lead, LeadEvent
from src.database.repository import LeadRepository
from src.verification.live_verifier import LiveVerifier


# Regex matching valid domain name (RFC compliant label structure)
DOMAIN_REGEX = re.compile(
    r"^(?:[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?\.)+[a-zA-Z]{2,63}$"
)


class DomainImporter:
    """Enterprise domain batch importer, cleaner, and live verification enqueuer."""

    @classmethod
    def clean_and_validate_domain(cls, raw_input: str) -> Optional[str]:
        """Validate domain, strip scheme, credentials, port, path, query, fragment, and whitespace.

        Returns lowercase domain if valid, or None if invalid.
        """
        if not raw_input or not isinstance(raw_input, str):
            return None

        raw = raw_input.strip().strip("\"' \t\r\n")
        if not raw:
            return None

        # Strip scheme if present
        if "://" in raw:
            raw = raw.split("://", 1)[1]
        elif raw.startswith("//"):
            raw = raw[2:]

        # Remove user info if any (e.g. user:pass@host)
        if "@" in raw:
            raw = raw.split("@", 1)[1]

        # Remove path, query params, fragments
        for char in ("/", "\\", "?", "#"):
            if char in raw:
                raw = raw.split(char, 1)[0]

        # Remove port if present
        if ":" in raw:
            raw = raw.split(":", 1)[0]

        domain = raw.lower().strip(".")
        if not domain or len(domain) > 253:
            return None

        # Check against regex
        if not DOMAIN_REGEX.match(domain):
            return None

        # Ensure labels don't start/end with hyphen
        labels = domain.split(".")
        for label in labels:
            if not label or label.startswith("-") or label.endswith("-"):
                return None

        return domain

    @classmethod
    def parse_from_text(cls, text: str) -> List[str]:
        """Parse raw lines or comma-separated domains from plain text."""
        raw_items = []
        for line in text.replace(",", "\n").splitlines():
            line = line.strip()
            if line and not line.startswith("#"):
                raw_items.append(line)
        return raw_items

    @classmethod
    def parse_from_csv(cls, content: Union[str, bytes]) -> List[str]:
        """Parse domain values from CSV content."""
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="replace")

        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if not rows:
            return []

        # Check header for target domain column
        header = [c.strip().lower() for c in rows[0]]
        target_col_idx = None
        for col_name in ("domain", "domain_name", "url", "website", "site", "host", "canonical_url"):
            if col_name in header:
                target_col_idx = header.index(col_name)
                break

        start_row = 1 if target_col_idx is not None else 0
        col_idx = target_col_idx if target_col_idx is not None else 0

        extracted = []
        for row in rows[start_row:]:
            if row and len(row) > col_idx:
                val = row[col_idx].strip()
                if val:
                    extracted.append(val)
        return extracted

    @classmethod
    def parse_from_xlsx(cls, content: bytes) -> List[str]:
        """Parse domain values from XLSX spreadsheet."""
        wb = load_workbook(filename=io.BytesIO(content), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        # Check header
        target_col_idx = None
        header_vals = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        for col_name in ("domain", "domain_name", "url", "website", "site", "host", "canonical_url"):
            if col_name in header_vals:
                target_col_idx = header_vals.index(col_name)
                break

        start_row = 1 if target_col_idx is not None else 0
        col_idx = target_col_idx if target_col_idx is not None else 0

        extracted = []
        for row in rows[start_row:]:
            if row and len(row) > col_idx and row[col_idx] is not None:
                val = str(row[col_idx]).strip()
                if val:
                    extracted.append(val)
        return extracted

    @classmethod
    def parse_from_json(cls, content: Union[str, bytes]) -> List[str]:
        """Parse domain values from JSON array or object."""
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="replace")

        data = json.loads(content)
        extracted = []

        if isinstance(data, list):
            for item in data:
                if isinstance(item, str):
                    extracted.append(item)
                elif isinstance(item, dict):
                    for k in ("domain", "domain_name", "url", "website", "site", "host", "canonical_url"):
                        if k in item and item[k]:
                            extracted.append(str(item[k]))
                            break
        elif isinstance(data, dict):
            for k in ("domains", "leads", "items", "data"):
                if k in data and isinstance(data[k], list):
                    return cls.parse_from_json(json.dumps(data[k]))
        return extracted

    @classmethod
    def extract_raw_entries(
        cls,
        content: Union[str, bytes, List[str]],
        format_hint: Optional[str] = None,
        filename: Optional[str] = None,
    ) -> List[str]:
        """Extract raw domain strings from any supported source format."""
        if isinstance(content, list):
            return [str(c) for c in content]

        # Determine format from filename or hint
        hint = (format_hint or "").lower()
        if not hint and filename:
            hint = filename.split(".")[-1].lower()

        if hint == "xlsx" and isinstance(content, bytes):
            return cls.parse_from_xlsx(content)
        elif hint == "csv":
            return cls.parse_from_csv(content)
        elif hint == "json":
            return cls.parse_from_json(content)
        elif hint in ("txt", "text"):
            text = content if isinstance(content, str) else content.decode("utf-8", errors="replace")
            return cls.parse_from_text(text)

        # Auto-detect if no hint
        if isinstance(content, bytes):
            if content.startswith(b"PK\x03\x04"):
                return cls.parse_from_xlsx(content)
            text = content.decode("utf-8", errors="replace")
        else:
            text = content

        text_stripped = text.strip()
        if text_stripped.startswith("[") or text_stripped.startswith("{"):
            try:
                return cls.parse_from_json(text)
            except Exception:
                pass

        if "\n" in text or "," in text:
            try:
                return cls.parse_from_csv(text)
            except Exception:
                pass

        return cls.parse_from_text(text)

    @classmethod
    def import_domains(
        cls,
        content: Union[str, bytes, List[str]],
        format_hint: Optional[str] = None,
        filename: Optional[str] = None,
        enqueue_verification: bool = True,
        session: Optional[Session] = None,
    ) -> Dict[str, Any]:
        """Execute full import pipeline:

        1. Extract raw entries.
        2. Clean and validate domains.
        3. Deduplicate against batch and database.
        4. Insert new leads.
        5. Enqueue for live verification.
        """
        raw_entries = cls.extract_raw_entries(content, format_hint=format_hint, filename=filename)

        valid_domains: List[str] = []
        batch_seen: Set[str] = set()
        batch_duplicates = 0
        invalid_entries: List[str] = []

        for entry in raw_entries:
            clean = cls.clean_and_validate_domain(entry)
            if clean:
                if clean in batch_seen:
                    batch_duplicates += 1
                else:
                    batch_seen.add(clean)
                    valid_domains.append(clean)
            else:
                invalid_entries.append(entry)

        # Deduplicate against database leads
        def _process_db(s: Session) -> Tuple[List[str], List[str]]:
            if not valid_domains:
                return [], []
            existing_db = set(
                d[0] for d in s.query(Lead.domain).filter(Lead.domain.in_(valid_domains)).all()
            )
            new_leads = [d for d in valid_domains if d not in existing_db]
            skipped = [d for d in valid_domains if d in existing_db]

            # Upsert new lead stubs
            for domain in new_leads:
                LeadRepository.upsert_lead(
                    session=s,
                    domain=domain,
                    canonical_url=f"https://{domain}",
                    business_name=domain,
                    status="UNKNOWN",
                    source="IMPORT",
                )
            return new_leads, skipped

        if session is not None:
            new_domains, db_duplicates = _process_db(session)
        else:
            with get_db() as s:
                new_domains, db_duplicates = _process_db(s)

        # Enqueue for live verification if requested
        enqueued_count = 0
        if enqueue_verification and new_domains:
            enqueued_count = len(new_domains)
            try:
                loop = asyncio.get_running_loop()
                loop.create_task(cls._verify_domains_worker(new_domains))
            except RuntimeError:
                pass

        return {
            "total_parsed": len(raw_entries),
            "valid_count": len(valid_domains),
            "unique_new_count": len(new_domains),
            "duplicate_count": len(db_duplicates) + batch_duplicates,
            "invalid_count": len(invalid_entries),
            "enqueued_count": enqueued_count,
            "imported_domains": new_domains,
            "skipped_duplicates": db_duplicates,
            "invalid_entries": invalid_entries,
        }

    @classmethod
    async def _verify_domains_worker(cls, domains: List[str]) -> None:
        """Background worker verifying enqueued domains."""
        from src.jobs.manager import JobManager
        from src.fingerprint.engine import FingerprintEngine
        from src.crawler.contact_extractor import ContactExtractor
        from src.scoring.lead_scorer import LeadScorer

        for domain in domains:
            try:
                v_res = await LiveVerifier.verify(domain, timeout=8.0)
                primary_tech = "Web Standard"
                tech_confidence = 0.3
                detected_techs = []
                evidence_list = []
                contacts = {"emails": [], "phones": [], "socials": {}}

                if v_res.is_live:
                    detections = FingerprintEngine.detect(
                        html=v_res.html_body,
                        headers=v_res.headers,
                        cookies=v_res.cookies,
                    )
                    detected_techs = [d.to_dict() for d in detections]
                    if detections:
                        top = detections[0]
                        primary_tech = top.name
                        tech_confidence = top.confidence
                        evidence_list = top.evidence

                    contacts = ContactExtractor.extract_from_html(v_res.html_body, v_res.canonical_url)

                score_res = LeadScorer.calculate(
                    has_tech_match=tech_confidence >= 0.5,
                    tech_confidence=tech_confidence,
                    target_country_match=False,
                    target_industry_match=False,
                    has_public_email=len(contacts.get("emails", [])) > 0,
                    has_public_phone=len(contacts.get("phones", [])) > 0,
                    has_ssl=v_res.has_ssl,
                    response_time_ms=v_res.response_time_ms,
                    is_live=v_res.is_live,
                    freshness_status=v_res.freshness_status,
                    tech_name=primary_tech,
                )

                with get_db() as session:
                    LeadRepository.upsert_lead(
                        session=session,
                        domain=domain,
                        canonical_url=v_res.canonical_url,
                        business_name=contacts.get("business_name") or domain,
                        description=contacts.get("description"),
                        status="LIVE" if v_res.is_live else "OFFLINE",
                        http_status=v_res.http_status,
                        has_ssl=v_res.has_ssl,
                        response_time_ms=v_res.response_time_ms,
                        primary_technology=primary_tech,
                        technology_confidence=tech_confidence,
                        evidence=evidence_list,
                        technologies=detected_techs,
                        emails=contacts.get("emails", []),
                        phones=contacts.get("phones", []),
                        socials=contacts.get("socials", {}),
                        lead_score=score_res.score,
                        score_label=score_res.label,
                        score_reasons=score_res.reasons,
                        source="IMPORT",
                    )
            except Exception:
                pass
