"""Comprehensive Test Suite for LeadForge SaaS Platform & Enterprise Features."""

import io
import json
import hmac
import hashlib
from openpyxl import Workbook, load_workbook
import pytest
from fastapi.testclient import TestClient

from src.api.app import app
from src.database.db import init_db, get_db
from src.database.models import Lead, Project, LeadList, Webhook, WebhookDelivery
from src.export.exporter import EnterpriseExporter
from src.export.async_exporter import AsyncExportManager
from src.importer.domain_importer import DomainImporter
from src.search.nl_parser import NLSearchParser
from src.webhooks.dispatcher import WebhookDispatcher


client = TestClient(app)


def setup_module():
    """Initialize database and seed baseline state."""
    init_db()


@pytest.fixture(autouse=True)
def cleanup_test_leads():
    """Ensure test leads are cleaned up before each test runs for isolation."""
    test_domains = [
        "already-exists.com",
        "brand-new-domain.com",
        "second-new-store.co.uk",
        "csv-store1.com",
        "csv-store2.de",
        "json-store1.org",
        "json-store2.net",
        "api-import-1.com",
        "api-import-2.com",
        "brand-new-domain-xyz.com",
        "second-new-store-xyz.co.uk",
    ]
    with get_db() as session:
        session.query(Lead).filter(Lead.domain.in_(test_domains)).delete(synchronize_session=False)
    yield


# ============================================================================
# 1. FORMULA INJECTION SANITIZATION (CSV / XLSX)
# ============================================================================

MALICIOUS_LEADS = [
    {
        "id": 101,
        "business_name": "=cmd|' /C calc'!A0",
        "domain": "+evil-domain.com",
        "canonical_url": "https://evil-domain.com",
        "primary_technology": "-DDE('param')",
        "technology_confidence": 0.99,
        "status": "LIVE",
        "http_status": 200,
        "has_ssl": True,
        "country": "@ATTACK_VECTOR",
        "industry": "=HYPERLINK('http://evil.com','Click')",
        "primary_email": "+admin@evil.com",
        "primary_phone": "-1234567890",
        "lead_score": 90,
        "score_label": "HOT",
        "score_reasons": ["=MALICIOUS_REASON"],
        "last_verified_at": "2026-09-07T12:00:00Z",
    }
]


def test_sanitize_spreadsheet_value_helper():
    """Verify formula injection sanitization function prefixes dangerous strings."""
    assert EnterpriseExporter.sanitize_formula_injection("=1+1") == "'=1+1"
    assert EnterpriseExporter.sanitize_formula_injection("+SUM(A1:A5)") == "'+SUM(A1:A5)"
    assert EnterpriseExporter.sanitize_formula_injection("-2+3") == "'-2+3"
    assert EnterpriseExporter.sanitize_formula_injection("@SUM") == "'@SUM"
    # Safe values remain untouched
    assert EnterpriseExporter.sanitize_formula_injection("Shopify Store") == "Shopify Store"
    assert EnterpriseExporter.sanitize_formula_injection(100) == 100
    assert EnterpriseExporter.sanitize_formula_injection(True) is True
    assert EnterpriseExporter.sanitize_formula_injection(None) is None


def test_csv_formula_injection_sanitization():
    """Verify CSV export sanitizes fields starting with =, +, -, @ by prepending single quote."""
    csv_output = EnterpriseExporter.to_csv(MALICIOUS_LEADS)

    # Values must be escaped with prepended single quote
    assert "'=cmd|' /C calc'!A0" in csv_output
    assert "'+evil-domain.com" in csv_output
    assert "'-DDE('param')" in csv_output
    assert "'@ATTACK_VECTOR" in csv_output
    assert "'+admin@evil.com" in csv_output

    # Raw unescaped dangerous formulas should not exist at column boundaries
    assert ",=cmd|" not in csv_output
    assert ",+evil-domain.com" not in csv_output
    assert ",@ATTACK_VECTOR" not in csv_output


def test_xlsx_formula_injection_sanitization():
    """Verify XLSX export sanitizes cells starting with =, +, -, @."""
    xlsx_bytes = EnterpriseExporter.to_xlsx(MALICIOUS_LEADS)
    assert len(xlsx_bytes) > 0
    assert xlsx_bytes[:4] == b"PK\x03\x04"

    wb = load_workbook(filename=io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb.active

    # Check cell values in row 2
    row_values = [cell.value for cell in ws[2]]
    assert "'=cmd|' /C calc'!A0" in row_values
    assert "'+evil-domain.com" in row_values
    assert "'-DDE('param')" in row_values
    assert "'@ATTACK_VECTOR" in row_values
    assert "'+admin@evil.com" in row_values


# ============================================================================
# 2. NATURAL LANGUAGE SEARCH PARSER ACCURACY
# ============================================================================

def test_nl_search_parser_opencart_uk():
    """Query: 'Find live OpenCart stores in the UK with public emails'"""
    q = "Find live OpenCart stores in the UK with public emails"
    parsed = NLSearchParser.parse(q)

    assert parsed.technology == "OpenCart"
    assert parsed.country == "United Kingdom"
    assert parsed.industry == "E-commerce"
    assert parsed.status == "LIVE"
    assert parsed.has_email is True


def test_nl_search_parser_shopify_germany():
    """Query: 'Shopify stores in Germany'"""
    q = "Shopify stores in Germany"
    parsed = NLSearchParser.parse(q)

    assert parsed.technology == "Shopify"
    assert parsed.country == "Germany"
    assert parsed.industry == "E-commerce"


def test_nl_search_parser_variations():
    """Verify various filter combinations and synonyms."""
    # Magento offline in France without email
    res1 = NLSearchParser.parse("offline Magento shops in France without email")
    assert res1.technology == "Magento"
    assert res1.country == "France"
    assert res1.status == "OFFLINE"
    assert res1.has_email is False
    assert res1.industry == "E-commerce"

    # WooCommerce fashion stores in the US with score > 80
    res2 = NLSearchParser.parse("WooCommerce fashion apparel in the US score > 80")
    assert res2.technology == "WooCommerce"
    assert res2.country == "United States"
    assert res2.industry == "Fashion"
    assert res2.min_score == 80


def test_api_search_parse_endpoints():
    """Verify /api/v1/search/parse GET and POST endpoints."""
    # GET
    res_get = client.get("/api/v1/search/parse?q=Find+live+OpenCart+stores+in+the+UK+with+public+emails")
    assert res_get.status_code == 200
    data_get = res_get.json()
    assert data_get["technology"] == "OpenCart"
    assert data_get["country"] == "United Kingdom"
    assert data_get["status"] == "LIVE"
    assert data_get["has_email"] is True

    # POST
    res_post = client.post("/api/v1/search/parse", json={"query": "Shopify stores in Germany"})
    assert res_post.status_code == 200
    data_post = res_post.json()
    assert data_post["technology"] == "Shopify"
    assert data_post["country"] == "Germany"


# ============================================================================
# 3. WEBHOOK HMAC-SHA256 SIGNATURE GENERATION
# ============================================================================

def test_webhook_hmac_signature_generation():
    """Verify deterministic HMAC-SHA256 signature signing and verification."""
    secret = "leadforge_enterprise_secret_key_2026"
    payload = json.dumps({"event": "lead.created", "lead_id": 42}, separators=(",", ":"))

    # Expected HMAC-SHA256
    expected_sig = hmac.new(
        secret.encode("utf-8"), payload.encode("utf-8"), hashlib.sha256
    ).hexdigest()

    generated_sig = WebhookDispatcher.sign_payload(payload, secret)
    assert generated_sig == expected_sig
    assert len(generated_sig) == 64  # 256 bits in hex

    # Verify signature verification method
    assert WebhookDispatcher.verify_signature(payload, secret, generated_sig) is True
    assert WebhookDispatcher.verify_signature(payload, "wrong_secret", generated_sig) is False
    assert WebhookDispatcher.verify_signature(payload + "tampered", secret, generated_sig) is False


def test_api_webhooks_crud_and_ping():
    """Verify Webhook CRUD operations and test ping endpoint."""
    # 1. Create Webhook
    create_res = client.post(
        "/api/v1/webhooks",
        json={
            "url": "https://webhook.site/leadforge-test",
            "events": ["lead.created", "export.completed"],
            "secret": "test_secret_key_123",
        },
    )
    assert create_res.status_code == 200
    webhook_data = create_res.json()
    webhook_id = webhook_data["id"]
    assert webhook_data["url"] == "https://webhook.site/leadforge-test"
    assert "lead.created" in webhook_data["events"]

    # 2. Get Webhook
    get_res = client.get(f"/api/v1/webhooks/{webhook_id}")
    assert get_res.status_code == 200
    assert get_res.json()["id"] == webhook_id

    # 3. Update Webhook
    update_res = client.put(
        f"/api/v1/webhooks/{webhook_id}",
        json={"is_active": False},
    )
    assert update_res.status_code == 200
    assert update_res.json()["is_active"] is False

    # 4. Reactivate and test ping
    client.put(f"/api/v1/webhooks/{webhook_id}", json={"is_active": True})
    ping_res = client.post(f"/api/v1/webhooks/{webhook_id}/ping")
    assert ping_res.status_code == 200
    assert ping_res.json()["status"] == "PING_DISPATCHED"

    # 5. Check deliveries recorded in database
    deliv_res = client.get(f"/api/v1/webhooks/{webhook_id}/deliveries")
    assert deliv_res.status_code == 200
    deliveries = deliv_res.json()
    assert len(deliveries) >= 1
    assert deliveries[0]["event_type"] == "test.ping"

    # 6. Delete Webhook
    del_res = client.delete(f"/api/v1/webhooks/{webhook_id}")
    assert del_res.status_code == 200
    assert del_res.json()["status"] == "DELETED"


# ============================================================================
# 4. IMPORT DOMAIN VALIDATION AND DEDUPLICATION
# ============================================================================

def test_domain_cleaning_and_validation():
    """Verify domain URL stripping (schemes, paths, queries, ports) and validation."""
    clean = DomainImporter.clean_and_validate_domain

    # Valid URLs and domains
    assert clean("https://store.example.com/checkout?step=1#cart") == "store.example.com"
    assert clean("http://www.myshop.co.uk:8080/products") == "www.myshop.co.uk"
    assert clean("ftp://user:pass@vendor.de/catalog") == "vendor.de"
    assert clean("  pure-domain.org  ") == "pure-domain.org"
    assert clean("//cdn.sub.domain.com/asset.js") == "cdn.sub.domain.com"

    # Invalid entries
    assert clean("invalid..domain.com") is None
    assert clean("http:///broken") is None
    assert clean("") is None
    assert clean("   ") is None
    assert clean("not-a-domain") is None
    assert clean("192.168.1.1") is None  # Pure IP not a named domain


def test_import_domains_from_text_and_deduplication():
    """Verify deduplication within batch and against existing DB leads."""
    # Seed an existing lead in DB
    with get_db() as session:
        existing = session.query(Lead).filter(Lead.domain == "already-exists.com").first()
        if not existing:
            lead = Lead(
                domain="already-exists.com",
                canonical_url="https://already-exists.com",
                business_name="Already Exists Store",
                status="LIVE",
                source="SEED",
            )
            session.add(lead)

    input_text = """
    https://already-exists.com/store
    ALREADY-EXISTS.COM
    https://brand-new-domain.com/shop
    BRAND-NEW-DOMAIN.COM
    http://second-new-store.co.uk:443/home?q=test
    invalid..domain..com
    not_a_valid_domain
    """

    result = DomainImporter.import_domains(
        content=input_text,
        format_hint="txt",
        enqueue_verification=False,
    )

    assert result["unique_new_count"] == 2
    assert "brand-new-domain.com" in result["imported_domains"]
    assert "second-new-store.co.uk" in result["imported_domains"]
    assert "already-exists.com" in result["skipped_duplicates"]
    assert result["invalid_count"] >= 2


def test_import_domains_from_csv():
    """Verify CSV import detecting domain column."""
    csv_content = """domain,business_name,notes
https://csv-store1.com/shop,Store 1,note
https://csv-store2.de/cart,Store 2,note
"""
    result = DomainImporter.import_domains(
        content=csv_content,
        format_hint="csv",
        enqueue_verification=False,
    )
    assert result["unique_new_count"] >= 2
    assert "csv-store1.com" in result["imported_domains"]
    assert "csv-store2.de" in result["imported_domains"]


def test_import_domains_from_json():
    """Verify JSON import supporting list of objects."""
    json_data = json.dumps([
        {"url": "https://json-store1.org/home"},
        {"domain": "json-store2.net"},
    ])
    result = DomainImporter.import_domains(
        content=json_data,
        format_hint="json",
        enqueue_verification=False,
    )
    assert result["unique_new_count"] >= 2
    assert "json-store1.org" in result["imported_domains"]
    assert "json-store2.net" in result["imported_domains"]


def test_api_import_text_endpoint():
    """Verify /api/v1/import/text endpoint."""
    res = client.post(
        "/api/v1/import/text",
        json={
            "domains": [
                "https://api-import-1.com/store",
                "api-import-2.com",
            ],
            "enqueue_verification": False,
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert data["unique_new_count"] >= 2
    assert "api-import-1.com" in data["imported_domains"]
    assert "api-import-2.com" in data["imported_domains"]


# ============================================================================
# 5. PROJECTS AND LEAD LISTS CRUD API
# ============================================================================

def test_projects_and_lead_lists_workflow():
    """Verify Projects and Lead Lists management."""
    # 1. Create Project
    p_res = client.post(
        "/api/v1/projects",
        json={"name": "Q3 Enterprise Outreach", "description": "High-value prospects"},
    )
    assert p_res.status_code == 200
    p_data = p_res.json()
    project_id = p_data["id"]

    # 2. Create Lead List under Project
    l_res = client.post(
        "/api/v1/lists",
        json={"name": "UK Retailers", "project_id": project_id},
    )
    assert l_res.status_code == 200
    l_data = l_res.json()
    list_id = l_data["id"]

    # Seed test lead
    with get_db() as session:
        test_lead = Lead(
            domain="project-member-test.com",
            canonical_url="https://project-member-test.com",
            business_name="Project Member Test",
            status="LIVE",
            source="TEST",
        )
        session.add(test_lead)
        session.flush()
        target_lead_id = test_lead.id

    # 3. Add Lead to List
    add_res = client.post(
        f"/api/v1/lists/{list_id}/leads",
        json={"lead_id": target_lead_id},
    )
    assert add_res.status_code == 200
    assert add_res.json()["status"] == "SUCCESS"

    # 4. Get Leads in List
    leads_res = client.get(f"/api/v1/lists/{list_id}/leads")
    assert leads_res.status_code == 200
    assert leads_res.json()["total"] >= 1

    # 5. Cleanup List & Project
    assert client.delete(f"/api/v1/lists/{list_id}").status_code == 200
    assert client.delete(f"/api/v1/projects/{project_id}").status_code == 200


# ============================================================================
# 6. MULTI-FORMAT AND ASYNC EXPORT VERIFICATION
# ============================================================================

def test_markdown_and_pdf_exports():
    """Verify Markdown and PDF export generators."""
    sample = [
        {
            "id": 1,
            "business_name": "OpenCart Demo",
            "domain": "opencart-demo.com",
            "primary_technology": "OpenCart",
            "status": "LIVE",
            "country": "United Kingdom",
            "primary_email": "demo@opencart-demo.com",
            "lead_score": 92,
            "score_label": "HOT",
        }
    ]
    # Markdown
    md_str = EnterpriseExporter.to_markdown(sample)
    assert "| Business Name |" in md_str
    assert "OpenCart Demo" in md_str

    # PDF
    pdf_bytes = EnterpriseExporter.to_pdf(sample)
    assert len(pdf_bytes) > 500
    assert pdf_bytes[:4] == b"%PDF"


def test_async_export_job_execution():
    """Verify async export job creation and status tracking via ExportJob."""
    job_id = AsyncExportManager.start_export_job(
        export_format="CSV",
        filters={"limit": 5},
    )
    assert job_id.startswith("exp_")

    job = AsyncExportManager.get_export_job(job_id)
    assert job is not None
    assert job["id"] == job_id
    assert job["status"] in ("PENDING", "PROCESSING", "COMPLETED")
